from abc import abstractmethod, ABC
from contextlib import closing
from dataclasses import dataclass, field
from typing import Any

from openpyxl import load_workbook

from xml_builder import XMLBuilder


@dataclass(frozen=True)
class SchemaField:
    name: str
    required: bool = True

@dataclass(frozen=True)
class SchemaEntity:
    schema: Schema
    fields: dict[str,Any]

    def has_value(self, field: SchemaField) -> bool:
        return self.fields.__contains__(field.name)

    def get_value(self, field: SchemaField):
        return self.fields.get(field.name)

    ## str value
    def set_value(self, field: SchemaField, value:str):
        self.fields[field.name] = value

    ## array value
    def append_value(self, field: SchemaField, value:str):
        if not self.has_value(field):
            self.fields[field.name] = []

        self.fields[field.name].append(value)

    def validate_matching(self, record, row_num):
        for field in self.schema.get_required_field_names():
            value = record.get(field)

            # Blank value is allowed
            if not Schema.has_value(value):
                continue

            current_value = getattr(self, field)

            if str(value) != str(current_value):
                raise ValueError(
                    f"Row {row_num}: {field} has value '{value}' "
                    f"but current bundle has '{current_value}'"
                )


@dataclass(frozen=True)
class Schema(ABC):
    name: str
    fields: list[SchemaField]
    entities: list[Any] = field(default_factory=list)

    def process_work_book(self, filename) -> Schema:
        with closing(load_workbook(filename, read_only=True)) as work_book:
            return self.process_worksheet(work_book.active)

    @abstractmethod
    def process_worksheet(self, work_sheet) -> Schema:
        pass

    @abstractmethod
    def process_entities(self) -> XMLBuilder:
        pass

    @staticmethod
    def row_is_empty(row):
        return all(cell is None or str(cell).strip() == "" for cell in row)

    @staticmethod
    def has_value(value):
        return value is not None and str(value).strip() != ""

    def get_required_field_names(self):
        return [f.name for f in self.fields if f.required]

    def get_field_names(self):
        return [f.name for f in self.fields]

    def validate_sheet(self, sheet):

        raw_headers = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))

        parsed_headers = []

        for col_num, header in enumerate(raw_headers, start=1):
            if header is None or str(header).strip() == "":
                raise ValueError(f"Unnamed header found in column {col_num}")

            parsed_headers.append(str(header).strip())

        # Check for duplicates
        duplicates = {h for h in parsed_headers if parsed_headers.count(h) > 1}
        if duplicates:
            raise ValueError(f"Duplicate headers found: {sorted(duplicates)}")

        # Check expected vs actual
        actual_fields = set(parsed_headers)

        all_fields = set(self.get_field_names())

        missing = all_fields - actual_fields
        if missing:
            raise ValueError(f"Missing columns: {sorted(missing)}")

        extra = actual_fields - all_fields
        if extra:
            raise ValueError(f"Unexpected columns: {sorted(extra)}")

        return parsed_headers

    def validate_required_fields(self, record, row_num):
        # print(type(record))
        # print(record)
        missing = [
            field
            for field in self.get_required_field_names()
            if not Schema.has_value(record.get(field))
        ]

        if missing:
            raise ValueError(f"Missing required field(s) at row {row_num}: {', '.join(missing)}")

    def create_entity_with_required_fields(self, record, row_num) -> SchemaEntity:
        self.validate_required_fields(record, row_num)

        create_dict = lambda keys: {k: record[k] for k in keys }

        return SchemaEntity(
            schema=self,
            fields = create_dict(self.get_required_field_names()))

    def create_entity_with_all_fields(self, record, row_num) -> SchemaEntity:
        self.validate_required_fields(record, row_num)

        create_dict = lambda keys: {k: record[k] for k in keys }

        return SchemaEntity(
            schema=self,
            fields = create_dict(self.get_field_names()))